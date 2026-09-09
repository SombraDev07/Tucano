"""Verifica que o .xlsx escrito pelo Tucano e lido por outra implementacao.

Round-trip proprio nao prova nada: um leitor e um escritor com o mesmo
mal-entendido concordam entre si. O que vale e outra implementacao concordar —
aqui, o openpyxl, que le o mesmo Office Open XML que o Excel abre.

    pixi run -e fixtures interop-xlsx
"""

import datetime as dt
import sys
from pathlib import Path

from openpyxl import load_workbook

ESCRITOS = Path("/tmp")


def valores(caminho, aba_esperada):
    wb = load_workbook(caminho, data_only=True)
    if wb.sheetnames != [aba_esperada]:
        raise AssertionError(f"abas {wb.sheetnames} != ['{aba_esperada}']")
    ws = wb[aba_esperada]
    return [list(linha) for linha in ws.iter_rows(values_only=True)]


def main():
    problemas = []

    linhas = valores(ESCRITOS / "tucano_xlsx_misto.xlsx", "Vendas")
    esperado = [
        ("nome", "qtd", "preco", "ok"),
        ("ana", 1, 1.5, True),
        ("bru & co", 20, 2.25, False),
        ("<carlos>", -3, -0.75, True),
        (None, 400, 100.0, False),
    ]
    for i, (obtido, quero) in enumerate(zip(linhas, esperado)):
        if tuple(obtido) != quero:
            problemas.append(f"misto linha {i}: {tuple(obtido)} != {quero}")
    if len(linhas) != len(esperado):
        problemas.append(f"misto: {len(linhas)} linhas, esperado {len(esperado)}")

    linhas = valores(ESCRITOS / "tucano_xlsx_datas.xlsx", "Datas")
    esperado = [
        ("quando", "carimbo"),
        (dt.datetime(2024, 1, 15), dt.datetime(2024, 1, 15, 8, 30, 0)),
        (dt.datetime(2024, 2, 29), dt.datetime(2024, 2, 29, 23, 59, 59)),
        (dt.datetime(1969, 12, 31), dt.datetime(1969, 12, 31, 23, 59, 59)),
    ]
    for i, (obtido, quero) in enumerate(zip(linhas, esperado)):
        if tuple(obtido) != quero:
            problemas.append(f"datas linha {i}: {tuple(obtido)} != {quero}")

    if problemas:
        print("interop .xlsx FALHOU:")
        for p in problemas:
            print("  ", p)
        sys.exit(1)

    print("lendo com openpyxl os arquivos escritos pelo Tucano:")
    print("  ok    misto      4 linhas  4 colunas  texto, inteiro, real, logico")
    print("  ok    datas      3 linhas  2 colunas  data e datahora, com hora exata")
    print("interop .xlsx confirmada")


if __name__ == "__main__":
    main()
